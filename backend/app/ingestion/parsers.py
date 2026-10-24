"""Turn raw bytes into text lines + tables + metadata. Originals are never modified."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from email import policy
from email.parser import BytesParser
from pathlib import PurePosixPath

SUPPORTED_TYPES = {
    ".csv": "csv",
    ".tsv": "csv",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "txt",
    ".md": "md",
    ".markdown": "md",
    ".eml": "eml",
}

MIME_TYPES = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
    "md": "text/markdown",
    "eml": "message/rfc822",
}

MAX_TABLE_ROWS = 50_000


class UnsupportedFile(ValueError):
    pass


@dataclass
class ParsedTable:
    name: str
    columns: list[str]
    rows: list[list[str]]
    truncated: bool = False

    @property
    def row_count(self) -> int:
        return len(self.rows)


@dataclass
class ParsedDocument:
    file_type: str
    text: str
    tables: list[ParsedTable] = field(default_factory=list)
    meta: dict = field(default_factory=dict)
    title: str | None = None

    @property
    def lines(self) -> list[str]:
        return self.text.splitlines()


def file_type_for(filename: str) -> str:
    ext = PurePosixPath(filename.lower()).suffix
    if ext not in SUPPORTED_TYPES:
        raise UnsupportedFile(f"Format non pris en charge : {ext or 'sans extension'}")
    return SUPPORTED_TYPES[ext]


def decode_text(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        text = data.decode("utf-16")
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252", errors="replace")
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _table_text(table: ParsedTable, limit: int = 200) -> str:
    """A searchable text rendering of a table: header + first rows, one row per line."""
    lines = [f"[Table {table.name}] colonnes : " + ", ".join(table.columns)]
    for row in table.rows[:limit]:
        lines.append(" | ".join(f"{c}={v}" for c, v in zip(table.columns, row, strict=False) if v != ""))
    if table.row_count > limit:
        lines.append(f"… {table.row_count - limit} lignes supplémentaires")
    return "\n".join(lines)


def _dedupe_columns(columns: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for i, col in enumerate(columns):
        name = col.strip() or f"colonne_{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        out.append(name)
    return out


def parse_csv(data: bytes, filename: str) -> ParsedDocument:
    text = decode_text(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [r for r in reader]
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return ParsedDocument("csv", "", meta={"rows": 0, "columns": []})
    columns = _dedupe_columns(rows[0])
    body = [[_clean(c) for c in r] + [""] * (len(columns) - len(r)) for r in rows[1:]]
    body = [r[: len(columns)] for r in body]
    truncated = len(body) > MAX_TABLE_ROWS
    table = ParsedTable(PurePosixPath(filename).name, columns, body[:MAX_TABLE_ROWS], truncated)
    return ParsedDocument(
        "csv",
        _table_text(table),
        [table],
        meta={"rows": len(body), "columns": columns, "delimiter": getattr(dialect, "delimiter", ",")},
    )


def parse_xlsx(data: bytes, filename: str) -> ParsedDocument:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    tables = []
    for ws in wb.worksheets:
        values = [[_clean(c) for c in row] for row in ws.iter_rows(values_only=True)]
        values = [r for r in values if any(r)]
        if not values:
            continue
        width = max(len(r) for r in values)
        columns = _dedupe_columns((values[0] + [""] * width)[:width])
        body = [(r + [""] * width)[:width] for r in values[1:]]
        tables.append(ParsedTable(ws.title, columns, body[:MAX_TABLE_ROWS], len(body) > MAX_TABLE_ROWS))
    props = wb.properties
    meta = {
        "sheets": [t.name for t in tables],
        "rows": sum(t.row_count for t in tables),
        "author": props.creator,
        "modified": props.modified.isoformat() if props.modified else None,
    }
    if len(tables) == 1:
        meta["columns"] = tables[0].columns
    wb.close()
    return ParsedDocument("xlsx", "\n\n".join(_table_text(t) for t in tables), tables, meta, props.title)


def parse_docx(data: bytes, filename: str) -> ParsedDocument:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    title = None
    for p in document.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading") or style == "title":
            level = int(style[-1]) if style[-1:].isdigit() else 1
            lines.append("#" * level + " " + text)
            title = title or text
        else:
            lines.append(text)
    tables = []
    for i, t in enumerate(document.tables, start=1):
        rows = [[_clean(c.text) for c in row.cells] for row in t.rows]
        if len(rows) >= 2:
            tables.append(ParsedTable(f"Tableau {i}", _dedupe_columns(rows[0]), rows[1:]))
    text = "\n".join(lines)
    if tables:
        text += "\n\n" + "\n\n".join(_table_text(t) for t in tables)
    cp = document.core_properties
    meta = {
        "author": cp.author or None,
        "modified": cp.modified.isoformat() if cp.modified else None,
        "tables": len(tables),
    }
    return ParsedDocument("docx", text, tables, meta, cp.title or title)


def parse_pdf(data: bytes, filename: str) -> ParsedDocument:
    import pdfplumber

    lines: list[str] = []
    tables = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        meta = {"pages": len(pdf.pages), "author": (pdf.metadata or {}).get("Author")}
        for page_no, page in enumerate(pdf.pages, start=1):
            page_text = page.extract_text() or ""
            lines.append(f"[Page {page_no}]")
            lines.extend(line for line in page_text.splitlines() if line.strip())
            for t_index, raw in enumerate(page.extract_tables() or [], start=1):
                rows = [[_clean(c) for c in row] for row in raw if row]
                if len(rows) >= 2:
                    tables.append(
                        ParsedTable(f"Page {page_no} — tableau {t_index}", _dedupe_columns(rows[0]), rows[1:])
                    )
        title = (pdf.metadata or {}).get("Title")
    if not any(line.strip() and not line.startswith("[Page") for line in lines):
        meta["warning"] = "Aucun texte extractible (PDF scanné ?) — OCR non inclus dans le MVP"
    return ParsedDocument("pdf", "\n".join(lines), tables, meta, title)


_HEADER_RE = re.compile(r"^(From|To|Cc|Subject|Date|De|À|Objet)\s*:\s*(.*)$", re.I)
_HEADER_MAP = {"de": "from", "à": "to", "objet": "subject"}


def _email_meta_from_text(text: str) -> tuple[dict, str] | None:
    """Plain-text email exports ("From: … / Subject: …" header block)."""
    head: dict[str, str] = {}
    lines = text.split("\n")
    i = 0
    while i < len(lines) and lines[i].strip():
        m = _HEADER_RE.match(lines[i].strip())
        if not m:
            break
        key = m.group(1).lower()
        head[_HEADER_MAP.get(key, key)] = m.group(2).strip()
        i += 1
    if "subject" in head and ("from" in head or "date" in head):
        return head, "\n".join(lines[i:]).strip()
    return None


def parse_eml(data: bytes, filename: str) -> ParsedDocument:
    msg = BytesParser(policy=policy.default).parsebytes(data)
    body_part = msg.get_body(preferencelist=("plain", "html"))
    body = body_part.get_content() if body_part else ""
    if body_part is not None and body_part.get_content_type() == "text/html":
        body = re.sub(r"<[^>]+>", " ", body)
    meta = {
        "from": str(msg.get("From", "")),
        "to": str(msg.get("To", "")),
        "subject": str(msg.get("Subject", "")),
        "date": str(msg.get("Date", "")),
        "email": True,
        "attachments": [p.get_filename() for p in msg.iter_attachments() if p.get_filename()],
    }
    text = f"Subject: {meta['subject']}\nFrom: {meta['from']}\nDate: {meta['date']}\n\n{body.strip()}"
    return ParsedDocument("eml", text, meta=meta, title=meta["subject"] or None)


def parse_text(data: bytes, filename: str, file_type: str) -> ParsedDocument:
    text = decode_text(data).strip("﻿")
    meta: dict = {}
    title = None
    email = _email_meta_from_text(text)
    if email:
        head, _ = email
        meta.update(head)
        meta["email"] = True
        title = head.get("subject")
    else:
        heading = next((line for line in text.split("\n") if line.startswith("# ")), None)
        title = heading[2:].strip() if heading else None
        # Key/value metadata often found in procedures and reports ("Version : 2.3").
        for key, label in (("version", "Version"), ("updated", "Dernière mise à jour"), ("date", "Date")):
            m = re.search(rf"^{label}\s*:\s*(.+)$", text, re.M)
            if m:
                meta[key] = m.group(1).strip()
    tables = _markdown_tables(text) if file_type == "md" else []
    return ParsedDocument(file_type, text, tables, meta, title)


def _markdown_tables(text: str) -> list[ParsedTable]:
    tables, block = [], []
    for line in text.split("\n") + [""]:
        if line.strip().startswith("|"):
            block.append(line)
            continue
        if len(block) >= 3 and re.match(r"^\|?\s*:?-{3,}", block[1].strip()):
            cells = [[c.strip() for c in row.strip().strip("|").split("|")] for row in block]
            tables.append(ParsedTable(f"Tableau {len(tables) + 1}", _dedupe_columns(cells[0]), cells[2:]))
        block = []
    return tables


def parse(data: bytes, filename: str) -> ParsedDocument:
    file_type = file_type_for(filename)
    if file_type == "csv":
        return parse_csv(data, filename)
    if file_type == "xlsx":
        return parse_xlsx(data, filename)
    if file_type == "docx":
        return parse_docx(data, filename)
    if file_type == "pdf":
        return parse_pdf(data, filename)
    if file_type == "eml":
        return parse_eml(data, filename)
    return parse_text(data, filename, file_type)
