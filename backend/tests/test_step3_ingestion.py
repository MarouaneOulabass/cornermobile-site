from __future__ import annotations

import io
import uuid
from email.message import EmailMessage

import pytest
from sqlalchemy import select

from app.config import get_settings
from app.domain.enums import DocumentStatus
from app.ingestion.chunking import chunk_lines
from app.ingestion.classifier import classify
from app.ingestion.parsers import UnsupportedFile, parse
from app.models import DataTable, Document, Job
from app.services.ingestion import IncomingFile, register_files
from app.services.pipeline import import_folder_sync
from app.services.search import search
from app.services.storage import get_storage
from tests.conftest import login

DEMO = get_settings().demo_data_dir


# --- Parsers -------------------------------------------------------------------------------------


def test_csv_with_bom_and_semicolons():
    data = "﻿ref;qte;prix\nA;2;10,5\nB;3;7\n".encode()
    doc = parse(data, "stock.csv")
    table = doc.tables[0]
    assert table.columns == ["ref", "qte", "prix"]
    assert table.rows == [["A", "2", "10,5"], ["B", "3", "7"]]
    assert doc.meta["rows"] == 2


def test_xlsx_sheets_become_tables():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Devis"
    ws.append(["quote_id", "amount_mad", "manual_reentry"])
    ws.append(["Q-1", 1200, "yes"])
    ws.append(["Q-2", 800.0, "no"])
    buf = io.BytesIO()
    wb.save(buf)
    doc = parse(buf.getvalue(), "devis.xlsx")
    assert doc.file_type == "xlsx"
    assert doc.tables[0].name == "Devis"
    assert doc.tables[0].rows == [["Q-1", "1200", "yes"], ["Q-2", "800", "no"]]


def test_docx_paragraphs_headings_and_tables():
    import docx

    d = docx.Document()
    d.add_heading("Procédure achats", level=1)
    d.add_paragraph("Saisie dans Excel puis ressaisie dans l'ERP.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Étape", "Outil"
    t.cell(1, 0).text, t.cell(1, 1).text = "Validation", "Email"
    buf = io.BytesIO()
    d.save(buf)
    doc = parse(buf.getvalue(), "proc.docx")
    assert doc.title == "Procédure achats"
    assert "# Procédure achats" in doc.text and "ressaisie" in doc.text
    assert doc.tables[0].columns == ["Étape", "Outil"]


def test_pdf_text_extraction():
    from weasyprint import HTML

    pdf = HTML(string="<h1>Contrat de maintenance</h1><p>Pénalité de 3% en cas de retard.</p>").write_pdf()
    doc = parse(pdf, "contrat.pdf")
    assert doc.meta["pages"] == 1
    assert "Pénalité de 3%" in doc.text


def test_eml_and_plain_text_email_exports():
    msg = EmailMessage()
    msg["From"], msg["To"], msg["Subject"], msg["Date"] = (
        "a@x.ma",
        "b@x.ma",
        "Relance facture",
        "Mon, 1 Sep 2026",
    )
    msg.set_content("Le client demande une copie de la facture.")
    eml = parse(bytes(msg), "relance.eml")
    assert eml.meta["subject"] == "Relance facture" and "copie de la facture" in eml.text

    txt = parse(
        b"From: a@x.ma\nTo: b@x.ma\nSubject: SLA client\nDate: 2026-08-08\n\nOn approche du delai.", "e.txt"
    )
    assert txt.meta["email"] is True and txt.title == "SLA client" and txt.meta["date"] == "2026-08-08"


def test_markdown_metadata_and_tables():
    md = "# Procédure 1\n\nVersion : 2.3\nDernière mise à jour : 2026-01-12\n\n| a | b |\n|---|---|\n| 1 | 2 |\n"
    doc = parse(md.encode(), "p.md")
    assert doc.meta == {"version": "2.3", "updated": "2026-01-12"}
    assert doc.tables[0].rows == [["1", "2"]]


def test_unsupported_extension():
    with pytest.raises(UnsupportedFile):
        parse(b"\x00", "photo.heic")


def test_chunks_keep_exact_line_numbers():
    text = "\n\n" + "\n".join(f"ligne {i}" for i in range(3, 80))
    chunks = chunk_lines(text, max_lines=20)
    lines = text.split("\n")
    for c in chunks:
        assert c.text == "\n".join(lines[c.start_line - 1 : c.end_line])
    assert chunks[0].start_line == 3


def test_classifier_on_column_names():
    assert (
        classify("data/x.csv", None, "", ["invoice_id", "days_overdue", "amount_mad"]).category == "finance"
    )
    assert (
        classify("data/x.csv", None, "", ["nc_id", "root_cause", "repeat_within_90d"]).category == "qualite"
    )
    assert classify("contrat_01.md", "Contrat client 1", "Pénalité 2%").category == "contrats"


# --- Pipeline ------------------------------------------------------------------------------------


def test_demo_folder_is_fully_indexed_and_answer_key_excluded(db, seeded):
    mission = seeded["atlaspro-distribution"]
    job = import_folder_sync(db, mission, DEMO / "atlaspro-distribution")
    assert job.status == "done" and job.progress == 100
    docs = db.scalars(select(Document).where(Document.mission_id == mission.id)).all()
    assert len(docs) == 37
    assert all(d.status == DocumentStatus.indexed for d in docs)
    assert not any(d.filename in ("ground_truth.json", "manifest.json") for d in docs)
    categories = {d.path: d.category for d in docs}
    assert categories["data/quotes.csv"] == "ventes"
    assert categories["data/invoices.csv"] == "finance"
    assert categories["documents/contracts/contract_01.md"] == "contrats"
    assert categories["documents/procedures/procedure_03.md"] == "procedures"
    quotes = db.scalar(select(DataTable).join(Document).where(Document.path == "data/quotes.csv"))
    assert quotes.row_count == 160 and quotes.column_types["prep_minutes"] == "nombre"
    assert quotes.column_types["manual_reentry"] == "booleen"
    # The original is kept byte for byte.
    original = (DEMO / "atlaspro-distribution" / "data" / "quotes.csv").read_bytes()
    stored = db.scalar(select(Document).where(Document.path == "data/quotes.csv"))
    assert get_storage().get(stored.storage_key) == original


def test_duplicates_versions_and_errors(db, seeded):
    mission = seeded["darna-facility-services"]
    files = [
        IncomingFile("notes/a.txt", b"Premiere version"),
        IncomingFile("copie/a_copie.txt", b"Premiere version"),
        IncomingFile("photo.heic", b"\x00\x01"),
    ]
    job, docs = register_files(db, mission, files)
    from app.jobs import run_job

    run_job(job.id)
    db.expire_all()
    statuses = {
        d.path: d.status for d in db.scalars(select(Document).where(Document.mission_id == mission.id))
    }
    assert statuses == {"notes/a.txt": "indexed", "copie/a_copie.txt": "duplicate", "photo.heic": "error"}

    job2, _ = register_files(db, mission, [IncomingFile("notes/a.txt", b"Deuxieme version")])
    run_job(job2.id)
    db.expire_all()
    versions = db.scalars(
        select(Document)
        .where(Document.mission_id == mission.id, Document.path == "notes/a.txt")
        .order_by(Document.version)
    ).all()
    assert [(v.version, v.is_current) for v in versions] == [(1, False), (2, True)]
    assert versions[1].previous_version_id == versions[0].id


def test_import_of_120_files_is_async_with_progress(client, seeded, db):
    headers = login(client, "consultant@auditflow.demo")
    mission = seeded["maghrebpack-industries"]
    files = [
        ("files", (f"lot/rapport_{i:03d}.txt", f"Rapport {i}\nIncident ligne {i % 7}".encode(), "text/plain"))
        for i in range(120)
    ]
    resp = client.post(f"/api/missions/{mission.id}/documents", files=files, headers=headers)
    assert resp.status_code == 202, resp.text
    body = resp.json()
    assert body["received"] == 120 and body["queued"] == 120 and body["errors"] == 0
    job = client.get(f"/api/missions/{mission.id}/jobs/{body['job_id']}", headers=headers).json()
    assert job["status"] == "done" and job["done"] == 120 and job["progress"] == 100
    stats = client.get(f"/api/missions/{mission.id}/documents/stats", headers=headers).json()
    assert stats["by_status"] == {"indexed": 120}
    assert db.get(Job, uuid.UUID(body["job_id"])).total == 120


def test_folder_upload_keeps_relative_paths(client, seeded):
    headers = login(client, "consultant@auditflow.demo")
    mission = seeded["darna-facility-services"]
    files = [("files", ("a.md", b"# Contrat\nPenalite 2%", "text/markdown"))]
    resp = client.post(
        f"/api/missions/{mission.id}/documents",
        files=files,
        data={"paths": ["dossier/contrats/a.md"]},
        headers=headers,
    )
    assert resp.json()["documents"][0]["path"] == "dossier/contrats/a.md"


def test_hybrid_search_returns_citations(db, seeded, client):
    mission = seeded["atlaspro-distribution"]
    import_folder_sync(db, mission, DEMO / "atlaspro-distribution")
    hits = search(db, mission.id, "modèle de devis vérifier les prix à la main", limit=5)
    assert hits and hits[0].filename == "email_08.txt"
    doc = db.get(Document, hits[0].document_id)
    lines = doc.text.split("\n")
    assert hits[0].excerpt == "\n".join(lines[hits[0].start_line - 1 : hits[0].end_line])

    headers = login(client, "dirigeant@atlaspro.demo")
    api_hits = client.get(
        f"/api/missions/{mission.id}/search", params={"q": "pénalité SLA"}, headers=headers
    ).json()
    assert any(h["filename"].startswith("contract_") for h in api_hits[:6])


def test_document_endpoints_respect_roles_and_tenants(client, seeded, db):
    mission = seeded["atlaspro-distribution"]
    import_folder_sync(db, mission, DEMO / "atlaspro-distribution")
    doc = db.scalar(select(Document).where(Document.path == "data/quotes.csv"))
    table = doc.tables[0]

    executive = login(client, "dirigeant@atlaspro.demo")
    assert client.get(f"/api/missions/{mission.id}/documents", headers=executive).status_code == 200
    rows = client.get(
        f"/api/missions/{mission.id}/tables/{table.id}", params={"rows": "1,3"}, headers=executive
    ).json()
    assert rows["row_numbers"] == [1, 3] and rows["rows"][0][0] == "Q-0001"
    upload = client.post(
        f"/api/missions/{mission.id}/documents",
        files=[("files", ("x.txt", b"x", "text/plain"))],
        headers=executive,
    )
    assert upload.status_code == 403

    other = login(client, "dirigeant@darna.demo")
    darna = seeded["darna-facility-services"]
    assert client.get(f"/api/missions/{mission.id}/documents/{doc.id}", headers=other).status_code == 404
    # Even through their own mission URL, another tenant's document id is refused.
    assert client.get(f"/api/missions/{darna.id}/documents/{doc.id}", headers=other).status_code == 404
    assert client.get(f"/api/missions/{darna.id}/tables/{table.id}", headers=other).status_code == 404

    consultant = login(client, "consultant@auditflow.demo")
    lines = client.get(f"/api/missions/{mission.id}/documents/{doc.id}/lines", headers=consultant).json()
    assert lines["total"] > 100
    raw = client.get(f"/api/missions/{mission.id}/documents/{doc.id}/download", headers=consultant)
    assert raw.content == (DEMO / "atlaspro-distribution/data/quotes.csv").read_bytes()
    patched = client.patch(
        f"/api/missions/{mission.id}/documents/{doc.id}", json={"category": "finance"}, headers=consultant
    ).json()
    assert patched["classification_method"] == "manuel" and patched["category"] == "finance"
